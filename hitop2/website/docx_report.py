from io import BytesIO
from pathlib import Path

from django.conf import settings
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


BLUE = "DBE7F5"
LIGHT_GRAY = "F2F2F2"
BORDER_GRAY = "D9D9D9"
TEXT_GRAY = RGBColor(120, 120, 120)

DISCLAIMER = (
    "Este relatório é gerado automaticamente, sendo os seus resultados "
    "meramente indicativos. A discussão dos resultados e respetiva "
    "interpretação deve ser feita por profissionais qualificados, "
    "enquadrados nos manuais técnicos e integrados com outra informação "
    "clínica. Conteúdo confidencial."
)

EDITABLE_FIELDS = (
    (
        "Motivo da Avaliação / Encaminhamento",
        "Contextualizar brevemente o motivo da procura/encaminhamento.",
        3,
    ),
    (
        "Progressão e Impacto da Sintomatologia",
        "Curso temporal: início, evolução, fatores precipitantes, principais "
        "queixas e sintomas; autonomia, impacto funcional no dia-a-dia; "
        "tentativas prévias de tratamento/intervenção.",
        5,
    ),
    (
        "Finalidade da Avaliação",
        "Objetivo: diagnóstico, caracterização funcional, apoio escolar ou "
        "tomada de decisão clínica.",
        3,
    ),
    (
        "Circunstâncias em que decorreu a avaliação",
        "Local, duração, colaboração/motivação, fatores interferentes e "
        "adaptações.",
        3,
    ),
    (
        "Histórico Desenvolvimental",
        "Gravidez, parto, desenvolvimento motor e da linguagem, marcos "
        "desenvolvimentais, socialização, comportamento e percurso académico.",
        5,
    ),
    (
        "Histórico e Dinâmica Familiar",
        "Estrutura, relações, acontecimentos significativos, saúde mental na "
        "família, contexto socioeconómico, condições materiais e rede de apoio.",
        6,
    ),
    (
        "Histórico Laboral",
        "Ambiente laboral, histórico de absentismo e relações com colegas de "
        "trabalho.",
        3,
    ),
    (
        "Histórico Médico e Psiquiátrico",
        "Doenças, lesões, diagnósticos prévios, internamentos, medicação e "
        "tratamentos farmacológicos e não farmacológicos.",
        6,
    ),
    (
        "Enquadramento Cultural e Religioso",
        "Identidade cultural, imigração, sentido de pertença, crenças e "
        "prática religiosa/espiritual.",
        4,
    ),
    (
        "Estilo de Vida e Circunstâncias Atuais",
        "Stressores, condição económica, rotinas, exercício, sono, consumo de "
        "substâncias, eventos de vida, recursos e rede de apoio.",
        5,
    ),
    (
        "Observações Clínicas",
        "Discurso, pensamento, contacto ocular, perceção, humor, cognição, "
        "insight, motivação e comportamentos relevantes.",
        6,
    ),
    (
        "Resultados de Outros Instrumentos de Avaliação",
        "Outros questionários, testes neuropsicológicos e entrevistas.",
        4,
    ),
)


def _set_cell_fill(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_border(cell, color=BORDER_GRAY, size="8"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def _set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def _add_section_banner(document, text):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(7.1)
    cell = table.cell(0, 0)
    cell.width = Inches(7.1)
    _set_cell_fill(cell, BLUE)
    _set_cell_border(cell, BLUE)
    _set_cell_margins(cell, top=110, bottom=110)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    paragraph.paragraph_format.space_after = Pt(0)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_analysis(document, analysis):
    paragraphs = analysis.get("paragraphs", [])
    if not paragraphs:
        return
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    _set_cell_fill(cell, LIGHT_GRAY)
    _set_cell_border(cell, LIGHT_GRAY)
    _set_cell_margins(cell, top=140, bottom=140, start=160, end=160)
    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    for text in paragraphs:
        paragraph = cell.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(6)


def _chart_font(size, bold=False):
    candidates = (
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    )
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def _centered_text(draw, position, text, font, fill="#222222"):
    box = draw.textbbox((0, 0), text, font=font)
    width = box[2] - box[0]
    draw.text((position[0] - width / 2, position[1]), text, font=font, fill=fill)


def _gradient_color(percentile):
    stops = (
        (0, (230, 184, 184)),
        (5, (230, 184, 184)),
        (15, (244, 227, 181)),
        (50, (201, 230, 200)),
        (85, (244, 227, 181)),
        (95, (230, 184, 184)),
        (100, (230, 184, 184)),
    )
    for (start, start_color), (end, end_color) in zip(stops, stops[1:]):
        if start <= percentile <= end:
            ratio = 0 if start == end else (percentile - start) / (end - start)
            color = tuple(
                round(a + (b - a) * ratio)
                for a, b in zip(start_color, end_color)
            )
            return tuple(round(channel * 0.75 + 255 * 0.25) for channel in color)
    return stops[-1][1]


def _split_chart(chart, max_items=8):
    items = chart["items"]
    if not items:
        return [chart]

    fragments = []
    for start in range(0, len(items), max_items):
        show_header = start == 0
        first_row_y = 55 if show_header else 22
        fragment_items = []
        for index, item in enumerate(items[start:start + max_items]):
            fragment_items.append({**item, "y": first_row_y + index * 28})
        fragments.append({
            "items": fragment_items,
            "height": first_row_y + len(fragment_items) * 28,
            "show_header": show_header,
        })
    return fragments


def _render_chart_png(chart):
    scale = 2
    width = 1200
    height = max(chart["height"], 85)
    image = Image.new("RGB", (width * scale, height * scale), "white")
    draw = ImageDraw.Draw(image)
    font_13 = _chart_font(13 * scale)
    font_14 = _chart_font(14 * scale)
    font_15 = _chart_font(15 * scale)

    def sx(value):
        return round(value * scale)

    graph_left = 640
    graph_right = 1170
    graph_width = graph_right - graph_left
    marks = (
        (1, "P1", graph_left + 12),
        (5, "P5", graph_left + 35),
        (15, "P15", graph_left + 85),
        (50, None, None),
        (85, "P85", graph_right - 85),
        (95, "P95", graph_right - 28),
        (99, "P99", graph_right + 6),
    )

    show_header = chart.get("show_header", True)
    if show_header:
        _centered_text(draw, (sx(430), sx(0)), "VALOR BRUTO", font_13)
        _centered_text(draw, (sx(540), sx(0)), "PERCENTIL", font_13)

    for percentile, label, label_x in marks:
        x = graph_left + percentile / 100 * graph_width
        dash = sx(4)
        y = sx(35 if show_header else 0)
        while y < sx(height - 20):
            draw.line((sx(x), y, sx(x), min(y + dash, sx(height - 20))), fill="#cfcfcf", width=scale)
            y += dash * 2
        if show_header and label:
            _centered_text(draw, (sx(label_x), sx(7)), label, font_13)

    for item in chart["items"]:
        y = item["y"]
        draw.text((sx(20), sx(y - 14)), item["name"], font=font_14, fill="#222222")
        score = (
            f'{item["score"]:.2f}'
            if item["is_valid"] and item["score"] is not None
            else "Nulo"
        )
        percentile_text = (
            str(item["percentile"])
            if item["is_valid"] and item["percentile"] is not None
            else "Nulo"
        )
        _centered_text(draw, (sx(430), sx(y - 15)), score, font_15)
        _centered_text(draw, (sx(540), sx(y - 15)), percentile_text, font_15)

        for pixel in range(graph_width):
            percentile = pixel / graph_width * 100
            color = _gradient_color(percentile)
            draw.rectangle(
                (sx(graph_left + pixel), sx(y - 9), sx(graph_left + pixel + 1), sx(y - 1)),
                fill=color,
            )

        if item["is_valid"] and item["x"] is not None:
            x = sx(item["x"])
            cy = sx(y - 5)
            radius = sx(5)
            draw.ellipse(
                (x - radius, cy - radius, x + radius, cy + radius),
                fill="white",
                outline="#999999",
                width=sx(3),
            )

    output = BytesIO()
    image.save(output, format="PNG", optimize=True, dpi=(192, 192))
    output.seek(0)
    return output


def _add_chart(document, chart):
    for fragment in _split_chart(chart):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run().add_picture(
            _render_chart_png(fragment),
            width=Inches(7.05),
        )


def _add_disclaimer(document):
    paragraph = document.add_paragraph(DISCLAIMER)
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(12)
    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(70, 70, 70)


def _add_editable_field(document, title, guidance, lines):
    heading = document.add_table(rows=1, cols=1)
    heading.alignment = WD_TABLE_ALIGNMENT.CENTER
    heading.autofit = False
    heading.columns[0].width = Inches(7.1)
    heading_cell = heading.cell(0, 0)
    heading_cell.width = Inches(7.1)
    _set_cell_fill(heading_cell, BLUE)
    _set_cell_border(heading_cell, BLUE)
    _set_cell_margins(heading_cell, top=80, bottom=80)
    heading_run = heading_cell.paragraphs[0].add_run(title)
    heading_run.bold = True
    _keep_with_next(heading_cell.paragraphs[0])

    field = document.add_table(rows=1, cols=1)
    field.alignment = WD_TABLE_ALIGNMENT.CENTER
    field.autofit = False
    field.columns[0].width = Inches(7.1)
    cell = field.cell(0, 0)
    cell.width = Inches(7.1)
    _set_cell_border(cell)
    _set_cell_margins(cell, top=150, bottom=150, start=150, end=150)
    paragraph = cell.paragraphs[0]
    placeholder = paragraph.add_run("Clique aqui e escreva a informação clínica.")
    placeholder.italic = True
    placeholder.font.color.rgb = TEXT_GRAY
    for _ in range(lines - 1):
        cell.add_paragraph("")

    note = document.add_paragraph(f"Exemplo: {guidance}")
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(14)
    for run in note.runs:
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = TEXT_GRAY


def _configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size in (("Title", 24), ("Heading 1", 19), ("Heading 2", 15)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)


def build_report_docx(context):
    document = Document()
    _configure_document(document)
    report = context["report"]

    logo_path = Path(settings.BASE_DIR) / "static" / "images" / "hitop_logo.png"
    if logo_path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(logo_path), width=Inches(5.8))

    if context.get("is_test_environment"):
        test_banner = document.add_table(rows=1, cols=1)
        test_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
        test_cell = test_banner.cell(0, 0)
        _set_cell_fill(test_cell, "FFF4CC")
        _set_cell_border(test_cell, "C8A94F", "16")
        _set_cell_margins(test_cell, top=140, bottom=140, start=160, end=160)
        test_paragraph = test_cell.paragraphs[0]
        test_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        test_run = test_paragraph.add_run(
            "SIMULAÇÃO / TESTE — este relatório não corresponde a dados "
            "clínicos reais e esta aplicação não foi incorporada na base "
            "normativa."
        )
        test_run.bold = True
        document.add_paragraph().paragraph_format.space_after = Pt(2)

    title = document.add_paragraph("Relatório Clínico", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    professional = document.add_paragraph()
    if context.get("is_test_environment"):
        run = professional.add_run(
            "Administrador responsável pelo teste: "
            f'{report["professional_name"]}'
        )
    else:
        run = professional.add_run(
            f'{report["professional_name"]}, {report["professional_area"]}, '
            f'Nº {report["professional_license"]}, Referência Certificação HiTOP'
        )
    run.italic = True

    details = document.add_table(rows=5, cols=2)
    details.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row, (label, value) in zip(
        details.rows,
        (
            ("Nome", report["patient_name"]),
            ("Idade", report["age"]),
            ("Sexo biológico", report["sex"]),
            ("Género com que se identifica", report["gender"]),
            ("Escolaridade", report["education"]),
        ),
    ):
        for cell in row.cells:
            _set_cell_border(cell, "FFFFFF", "0")
            _set_cell_margins(cell, top=30, bottom=30)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        label_run.bold = True
        row.cells[1].paragraphs[0].add_run(str(value))

    date = report["submission_date"]
    date_value = date.strftime("%d/%m/%Y") if date else "-"
    paragraph = document.add_paragraph()
    paragraph.add_run("Data de Preenchimento: ").bold = True
    paragraph.add_run(date_value)

    intro = document.add_table(rows=1, cols=1)
    intro_cell = intro.cell(0, 0)
    _set_cell_fill(intro_cell, BLUE)
    _set_cell_border(intro_cell, BLUE)
    _set_cell_margins(intro_cell, top=180, bottom=180, start=180, end=180)
    intro_cell.paragraphs[0].add_run(
        "O HiTOP:SR (Hierarchical Taxonomy of Psychopathology - Self-Report) "
        "permite medir padrões de sintomas psicológicos segundo uma abordagem "
        "baseada na evidência científica. Os resultados são informativos e não "
        "têm valor de diagnóstico, devendo ser integrados com outras fontes de "
        "informação e acompanhamento profissional."
    )

    document.add_heading("Resultados", level=1)
    attention = context["attention_checks"]
    if attention["incorrect"]:
        warning = document.add_paragraph()
        warning.add_run("Atenção: ").bold = True
        warning.add_run(
            f'O paciente falhou {attention["incorrect"]} de '
            f'{attention["total"]} checks de atenção.'
        )

    _add_section_banner(document, "Perfil Global")
    document.add_paragraph(
        "O quadro abaixo apresenta a pontuação bruta e o percentil para cada "
        "dimensão avaliada. O percentil compara os valores individuais com os "
        "valores de referência."
    )
    _add_analysis(document, context["global_analysis"])
    _add_chart(document, context["global_chart_data"])
    _add_disclaimer(document)

    _add_section_banner(document, "Perfil Detalhado")
    for index, section_data in enumerate(context["detailed_sections"]):
        if index:
            document.add_page_break()
        heading = document.add_heading(section_data["title"], level=2)
        _keep_with_next(heading)
        _add_analysis(document, section_data["analysis"])
        _add_chart(document, section_data["chart"])
        _add_disclaimer(document)

    document.add_page_break()
    document.add_heading("Anamnese e Informação Clínica Complementar", level=1)
    document.add_paragraph(
        "Preencha os campos abaixo diretamente no Word. As caixas aumentam "
        "automaticamente à medida que o texto é introduzido."
    )
    for field in EDITABLE_FIELDS:
        _add_editable_field(document, *field)

    document.add_heading("Conclusão e Recomendações", level=1)
    _add_editable_field(
        document,
        "Síntese Clínica",
        "Síntese integrativa, intervenção aconselhada, frequência, objetivos "
        "iniciais, necessidades específicas e encaminhamentos.",
        8,
    )

    document.add_paragraph("\n\n")
    signature = document.add_table(rows=1, cols=2)
    signature.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in signature.row_cells(0):
        _set_cell_fill(cell, LIGHT_GRAY)
        _set_cell_border(cell, LIGHT_GRAY)
        _set_cell_margins(cell, top=200, bottom=200, start=180, end=180)
    signature.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    signature.cell(0, 0).paragraphs[0].add_run("____________________________\nASSINATURA")
    info = signature.cell(0, 1).paragraphs[0]
    name_run = info.add_run(str(report["professional_name"]).upper())
    name_run.bold = True
    if context.get("is_test_environment"):
        info.add_run("\nADMINISTRADOR — AMBIENTE DE TESTE")
    else:
        info.add_run(
            f'\n{str(report["professional_area"]).upper()}'
            f'\nNº CÉDULA PROFISSIONAL: {report["professional_license"]}'
            "\nNº CERTIFICAÇÃO HITOP"
        )
    _add_disclaimer(document)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
